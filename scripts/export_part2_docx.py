"""Export PART2_SCIO_CLONE_PLAN.md to a Word document (.docx) with embedded images.

Usage:
    pip install pypandoc-binary python-docx
    python scripts/export_part2_docx.py
"""

from pathlib import Path

import pypandoc
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
SOURCE_MD = ROOT / "PART2_SCIO_CLONE_PLAN.md"
OUTPUT_DOCX = ROOT / "docs" / "PART2_SCIO_CLONE_PLAN.docx"


def convert_markdown() -> None:
    """Convert Markdown to DOCX with pandoc; relative image paths resolve from repo root."""
    pypandoc.convert_file(
        str(SOURCE_MD),
        "docx",
        format="gfm",
        outputfile=str(OUTPUT_DOCX),
        extra_args=[f"--resource-path={ROOT}"],
    )


def set_cell_borders(table) -> None:
    """Pandoc tables have no borders by default; add thin grey grid lines."""
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "A0A0A0")
        borders.append(el)
    tbl_pr.append(borders)


def shade_header_row(table) -> None:
    """Give the first row a light background so tables are easy to scan."""
    for cell in table.rows[0].cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), "E8EEF7")
        tc_pr.append(shd)
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True


def polish_document() -> None:
    """Apply fonts, table styling and margins for a clean, readable document."""
    doc = Document(str(OUTPUT_DOCX))

    for style_name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        if style_name in doc.styles:
            style = doc.styles[style_name]
            style.font.name = "Calibri"
            style.font.size = Pt(10.5)

    for level, size in ((1, 18), (2, 15), (3, 13), (4, 12)):
        name = f"Heading {level}"
        if name in doc.styles:
            style = doc.styles[name]
            style.font.name = "Calibri"
            style.font.size = Pt(size)
            style.font.color.rgb = RGBColor(0x1F, 0x3A, 0x68)

    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_cell_borders(table)
        if table.rows:
            shade_header_row(table)
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(9.5)

    for section in doc.sections:
        section.left_margin = section.right_margin = Pt(54)
        section.top_margin = section.bottom_margin = Pt(54)

    doc.save(str(OUTPUT_DOCX))


def main() -> None:
    OUTPUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    convert_markdown()
    polish_document()
    print(f"Wrote {OUTPUT_DOCX} ({OUTPUT_DOCX.stat().st_size // 1024} KB)")


if __name__ == "__main__":
    main()
